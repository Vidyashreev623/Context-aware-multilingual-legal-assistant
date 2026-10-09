"""
DOCX text extraction (Phase 3).

Extracts raw text (paragraphs + table cells) from .docx files using
python-docx. Like extract_pdf.py, this does not clean/normalize text.

Usage:
    python scripts/extraction/extract_docx.py path/to/file.docx
"""
from __future__ import annotations

import sys
from pathlib import Path


class DocxExtractionError(RuntimeError):
    pass


def extract_docx_text(path: str | Path) -> str:
    path = Path(path)
    if not path.exists():
        raise DocxExtractionError(f"File not found: {path}")

    try:
        import docx
    except ImportError as exc:
        raise DocxExtractionError(
            "python-docx is not installed. Run: pip install -r backend/requirements.txt"
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise DocxExtractionError(f"Could not open DOCX '{path}': {exc}") from exc

    parts: list[str] = []

    # Paragraphs, in document order.
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    # Tables (rental agreements / contracts often carry key terms in tables).
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    if not text.strip():
        raise DocxExtractionError(f"No extractable text found in '{path}'.")
    return text


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python extract_docx.py <file.docx>", file=sys.stderr)
        sys.exit(1)
    try:
        text = extract_docx_text(sys.argv[1])
    except DocxExtractionError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
    print(f"Extracted {len(text)} chars.")
    print("--- first 500 chars ---")
    print(text[:500])
